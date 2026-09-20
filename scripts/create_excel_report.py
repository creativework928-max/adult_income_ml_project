"""
===============================================================================
ADULT INCOME ML PROJECT
Professional Excel Report Generator
===============================================================================

Purpose
-------
Generate a professional Excel report from the outputs produced by the
Task 1 data preprocessing pipeline.

Expected project structure
--------------------------
adult_income_ml_project/
│
├── data/
│   ├── processed/
│   │   ├── cleaned_dataset.csv
│   │   └── preprocessing_summary.csv
│   │
│   └── splits/
│       ├── train_raw.csv
│       ├── test_raw.csv
│       ├── y_train.csv
│       └── y_test.csv
│
├── outputs/
│   ├── figures/
│   └── reports/
│
├── models/
│   └── preprocessing_pipeline.joblib
│
├── src/
│   └── ...
│
└── scripts/
    └── create_excel_report.py

Output
------
outputs/reports/task_1_data_preprocessing_report.xlsx

Author
------
Adult Income ML Project
===============================================================================
"""

from __future__ import annotations

import sys
from pathlib import Path
from datetime import datetime

import numpy as np
import pandas as pd

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# =============================================================================
# PROJECT PATHS
# =============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

DATA_DIR = PROJECT_ROOT / "data"
PROCESSED_DIR = DATA_DIR / "processed"
SPLITS_DIR = DATA_DIR / "splits"

OUTPUT_DIR = PROJECT_ROOT / "outputs"
REPORTS_DIR = OUTPUT_DIR / "reports"

REPORTS_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = (
    REPORTS_DIR /
    "task_1_data_preprocessing_report.xlsx"
)


# =============================================================================
# DATA FILES
# =============================================================================

CLEANED_DATA_FILE = (
    PROCESSED_DIR /
    "cleaned_dataset.csv"
)

SUMMARY_FILE = (
    PROCESSED_DIR /
    "preprocessing_summary.csv"
)

TRAIN_FILE = (
    SPLITS_DIR /
    "train_raw.csv"
)

TEST_FILE = (
    SPLITS_DIR /
    "test_raw.csv"
)

Y_TRAIN_FILE = (
    SPLITS_DIR /
    "y_train.csv"
)

Y_TEST_FILE = (
    SPLITS_DIR /
    "y_test.csv"
)


# =============================================================================
# PROFESSIONAL COLOR PALETTE
# =============================================================================

NAVY = "0B1F33"
DARK_BLUE = "12355B"
BLUE = "2563EB"
LIGHT_BLUE = "DBEAFE"

CYAN = "06B6D4"
LIGHT_CYAN = "CFFAFE"

GREEN = "10B981"
LIGHT_GREEN = "D1FAE5"

ORANGE = "F59E0B"
LIGHT_ORANGE = "FEF3C7"

RED = "EF4444"
LIGHT_RED = "FEE2E2"

PURPLE = "7C3AED"
LIGHT_PURPLE = "EDE9FE"

SLATE = "64748B"
LIGHT_SLATE = "E2E8F0"

VERY_LIGHT = "F8FAFC"
WHITE = "FFFFFF"
BLACK = "111827"


# =============================================================================
# EXCEL STYLE OBJECTS
# =============================================================================

thin_gray = Side(
    style="thin",
    color="CBD5E1"
)

medium_blue = Side(
    style="medium",
    color=BLUE
)

border = Border(
    left=thin_gray,
    right=thin_gray,
    top=thin_gray,
    bottom=thin_gray,
)

header_fill = PatternFill(
    "solid",
    fgColor=NAVY
)

subheader_fill = PatternFill(
    "solid",
    fgColor=DARK_BLUE
)

light_blue_fill = PatternFill(
    "solid",
    fgColor=LIGHT_BLUE
)

light_green_fill = PatternFill(
    "solid",
    fgColor=LIGHT_GREEN
)

light_orange_fill = PatternFill(
    "solid",
    fgColor=LIGHT_ORANGE
)

light_red_fill = PatternFill(
    "solid",
    fgColor=LIGHT_RED
)

light_purple_fill = PatternFill(
    "solid",
    fgColor=LIGHT_PURPLE
)

light_gray_fill = PatternFill(
    "solid",
    fgColor=VERY_LIGHT
)


# =============================================================================
# HELPER FUNCTIONS
# =============================================================================

def safe_read_csv(path: Path) -> pd.DataFrame:
    """
    Read CSV if it exists.

    Raises a helpful error if the file is missing.
    """

    if not path.exists():
        raise FileNotFoundError(
            f"Required file was not found:\n{path}\n\n"
            "Run the preprocessing pipeline first:\n"
            "python run_preprocessing.py"
        )

    return pd.read_csv(path)


def apply_title(
    ws,
    title: str,
    subtitle: str | None = None,
    end_column: int = 8,
):
    """
    Add a professional title section.
    """

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=end_column,
    )

    cell = ws.cell(
        row=1,
        column=1,
        value=title,
    )

    cell.fill = PatternFill(
        "solid",
        fgColor=NAVY,
    )

    cell.font = Font(
        color=WHITE,
        bold=True,
        size=20,
    )

    cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    ws.row_dimensions[1].height = 34

    if subtitle:

        ws.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=end_column,
        )

        subtitle_cell = ws.cell(
            row=2,
            column=1,
            value=subtitle,
        )

        subtitle_cell.font = Font(
            color=SLATE,
            italic=True,
            size=10,
        )

        subtitle_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        ws.row_dimensions[2].height = 22


def style_header_row(ws, row_number: int):
    """
    Apply professional header formatting.
    """

    for cell in ws[row_number]:

        if cell.value is not None:

            cell.fill = header_fill

            cell.font = Font(
                color=WHITE,
                bold=True,
                size=10,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            cell.border = border

    ws.row_dimensions[row_number].height = 30


def style_data_range(
    ws,
    start_row: int,
    end_row: int,
    start_column: int,
    end_column: int,
):
    """
    Apply consistent formatting to a rectangular data range.
    """

    for row in ws.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=start_column,
        max_col=end_column,
    ):

        for cell in row:

            cell.border = border

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )


def auto_fit_columns(
    ws,
    min_width: int = 10,
    max_width: int = 42,
):
    """
    Automatically size columns based on cell content.
    """

    for column_cells in ws.columns:

        max_length = 0
        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            try:
                value_length = len(
                    str(cell.value)
                    if cell.value is not None
                    else ""
                )

                max_length = max(
                    max_length,
                    value_length,
                )

            except Exception:
                pass

        width = min(
            max(
                max_length + 2,
                min_width,
            ),
            max_width,
        )

        ws.column_dimensions[
            column_letter
        ].width = width


def add_excel_table(
    ws,
    start_row: int,
    end_row: int,
    start_column: int,
    end_column: int,
    table_name: str,
):
    """
    Add a styled Excel table.
    """

    if end_row <= start_row:
        return

    ref = (
        f"{get_column_letter(start_column)}"
        f"{start_row}:"
        f"{get_column_letter(end_column)}"
        f"{end_row}"
    )

    table = Table(
        displayName=table_name,
        ref=ref,
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    table.tableStyleInfo = style

    ws.add_table(table)


def add_kpi_card(
    ws,
    cell_range: str,
    label: str,
    value,
    fill_color: str,
):
    """
    Create a dashboard KPI card.
    """

    ws.merge_cells(cell_range)

    start_cell = ws[cell_range.split(":")[0]]

    start_cell.value = (
        f"{label}\n"
        f"{value}"
    )

    start_cell.fill = PatternFill(
        "solid",
        fgColor=fill_color,
    )

    start_cell.font = Font(
        color=WHITE,
        bold=True,
        size=13,
    )

    start_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    start_cell.border = Border(
        left=medium_blue,
        right=medium_blue,
        top=medium_blue,
        bottom=medium_blue,
    )


def normalize_target(value):
    """
    Normalize target labels into readable values.
    """

    if pd.isna(value):
        return "Missing"

    value = str(value).strip()

    value = value.replace(".", "")

    if value == "0":
        return "<=50K"

    if value == "1":
        return ">50K"

    return value


# =============================================================================
# LOAD DATA
# =============================================================================

def load_project_data():

    cleaned_df = safe_read_csv(
        CLEANED_DATA_FILE
    )

    summary_df = safe_read_csv(
        SUMMARY_FILE
    )

    train_df = safe_read_csv(
        TRAIN_FILE
    )

    test_df = safe_read_csv(
        TEST_FILE
    )

    y_train_df = (
        safe_read_csv(Y_TRAIN_FILE)
        if Y_TRAIN_FILE.exists()
        else None
    )

    y_test_df = (
        safe_read_csv(Y_TEST_FILE)
        if Y_TEST_FILE.exists()
        else None
    )

    return (
        cleaned_df,
        summary_df,
        train_df,
        test_df,
        y_train_df,
        y_test_df,
    )


# =============================================================================
# DASHBOARD SHEET
# =============================================================================

def create_dashboard(
    wb,
    cleaned_df,
    train_df,
    test_df,
    summary_df,
):
    """
    Create executive-level dashboard.
    """

    ws = wb.create_sheet(
        "Dashboard",
        0,
    )

    ws.sheet_view.showGridLines = False

    apply_title(
        ws,
        "TASK 1 — DATA PREPROCESSING DASHBOARD",
        (
            "Adult Census Income Dataset | "
            "Professional preprocessing summary"
        ),
        end_column=10,
    )

    # -------------------------------------------------------------------------
    # KPI VALUES
    # -------------------------------------------------------------------------

    total_rows = len(cleaned_df)

    total_features = (
        len(cleaned_df.columns) - 1
        if "income" in cleaned_df.columns
        else len(cleaned_df.columns)
    )

    train_rows = len(train_df)
    test_rows = len(test_df)

    missing_total = int(
        cleaned_df.isnull().sum().sum()
    )

    train_ratio = (
        train_rows / total_rows
        if total_rows
        else 0
    )

    test_ratio = (
        test_rows / total_rows
        if total_rows
        else 0
    )

    # KPI cards
    add_kpi_card(
        ws,
        "A4:B6",
        "TOTAL RECORDS",
        f"{total_rows:,}",
        BLUE,
    )

    add_kpi_card(
        ws,
        "C4:D6",
        "FEATURES",
        f"{total_features}",
        CYAN,
    )

    add_kpi_card(
        ws,
        "E4:F6",
        "TRAINING",
        f"{train_rows:,}",
        GREEN,
    )

    add_kpi_card(
        ws,
        "G4:H6",
        "TESTING",
        f"{test_rows:,}",
        PURPLE,
    )

    add_kpi_card(
        ws,
        "I4:J6",
        "MISSING VALUES",
        f"{missing_total:,}",
        ORANGE,
    )

    # -------------------------------------------------------------------------
    # PROJECT INFORMATION
    # -------------------------------------------------------------------------

    ws["A8"] = "Project Information"
    ws["A8"].fill = subheader_fill
    ws["A8"].font = Font(
        color=WHITE,
        bold=True,
        size=12,
    )

    info = [
        ("Dataset", "UCI Adult / Census Income"),
        ("Task", "Data Preprocessing"),
        ("Split Strategy", "80% Train / 20% Test"),
        ("Sampling", "Stratified"),
        ("Random State", 42),
        ("Numerical Imputation", "Median"),
        ("Categorical Imputation", "Most Frequent"),
        ("Outlier Strategy", "1st–99th Percentile Clipping"),
        ("Numerical Scaling", "StandardScaler"),
        ("Categorical Encoding", "OneHotEncoder"),
        ("Leakage Prevention", "Pipeline fitted on training data only"),
        (
            "Report Generated",
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ),
    ]

    start_row = 9

    for index, (label, value) in enumerate(info):

        row = start_row + index

        ws.cell(
            row=row,
            column=1,
            value=label,
        )

        ws.cell(
            row=row,
            column=2,
            value=value,
        )

        ws.cell(
            row=row,
            column=1,
        ).font = Font(
            bold=True,
            color=NAVY,
        )

        ws.cell(
            row=row,
            column=1,
        ).fill = light_gray_fill

        ws.cell(
            row=row,
            column=2,
        ).fill = PatternFill(
            "solid",
            fgColor=WHITE,
        )

        ws.cell(
            row=row,
            column=1,
        ).border = border

        ws.cell(
            row=row,
            column=2,
        ).border = border

    # -------------------------------------------------------------------------
    # TRAIN / TEST PIE CHART
    # -------------------------------------------------------------------------

    chart = PieChart()

    labels = ["Training", "Testing"]
    values = [train_rows, test_rows]

    chart_ws = wb.create_sheet(
        "_DashboardData"
    )

    chart_ws.sheet_state = "hidden"

    chart_ws["A1"] = "Split"
    chart_ws["B1"] = "Records"

    for i, (label, value) in enumerate(
        zip(labels, values),
        start=2,
    ):

        chart_ws.cell(
            row=i,
            column=1,
            value=label,
        )

        chart_ws.cell(
            row=i,
            column=2,
            value=value,
        )

    data = Reference(
        chart_ws,
        min_col=2,
        min_row=1,
        max_row=3,
    )

    cats = Reference(
        chart_ws,
        min_col=1,
        min_row=2,
        max_row=3,
    )

    chart.add_data(
        data,
        titles_from_data=True,
    )

    chart.set_categories(cats)

    chart.title = "Train / Test Split"

    chart.height = 7
    chart.width = 11

    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True

    ws.add_chart(
        chart,
        "D9",
    )

    # -------------------------------------------------------------------------
    # TARGET DISTRIBUTION
    # -------------------------------------------------------------------------

    if "income" in cleaned_df.columns:

        target_counts = (
            cleaned_df["income"]
            .map(normalize_target)
            .value_counts()
        )

        target_start_row = 22

        ws.cell(
            target_start_row,
            1,
            "Target Distribution",
        )

        ws.cell(
            target_start_row,
            1,
        ).fill = subheader_fill

        ws.cell(
            target_start_row,
            1,
        ).font = Font(
            color=WHITE,
            bold=True,
            size=12,
        )

        ws.cell(
            target_start_row + 1,
            1,
            "Income Class",
        )

        ws.cell(
            target_start_row + 1,
            2,
            "Records",
        )

        for i, (label, count) in enumerate(
            target_counts.items(),
            start=target_start_row + 2,
        ):

            ws.cell(i, 1, label)
            ws.cell(i, 2, int(count))

        style_header_row(
            ws,
            target_start_row + 1,
        )

        style_data_range(
            ws,
            target_start_row + 2,
            target_start_row + 1 + len(target_counts),
            1,
            2,
        )

        add_excel_table(
            ws,
            target_start_row + 1,
            target_start_row + 1 + len(target_counts),
            1,
            2,
            "TargetDistribution",
        )

    ws.freeze_panes = "A4"

    auto_fit_columns(ws)

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 48


# =============================================================================
# SUMMARY SHEET
# =============================================================================

def create_summary_sheet(
    wb,
    summary_df,
):
    """
    Create preprocessing summary sheet.
    """

    ws = wb.create_sheet(
        "Preprocessing Summary"
    )

    ws.sheet_view.showGridLines = False

    apply_title(
        ws,
        "PREPROCESSING SUMMARY",
        "Key metrics generated by the preprocessing pipeline",
        end_column=max(
            2,
            len(summary_df.columns),
        ),
    )

    start_row = 4

    for col_index, column in enumerate(
        summary_df.columns,
        start=1,
    ):

        cell = ws.cell(
            start_row,
            col_index,
            column,
        )

        cell.fill = header_fill
        cell.font = Font(
            color=WHITE,
            bold=True,
        )

    for row_index, row in enumerate(
        summary_df.itertuples(index=False),
        start=start_row + 1,
    ):

        for col_index, value in enumerate(
            row,
            start=1,
        ):

            cell = ws.cell(
                row_index,
                col_index,
                value,
            )

            cell.border = border

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

    style_header_row(
        ws,
        start_row,
    )

    add_excel_table(
        ws,
        start_row,
        start_row + len(summary_df),
        1,
        len(summary_df.columns),
        "PreprocessingSummary",
    )

    ws.freeze_panes = "A5"

    auto_fit_columns(ws)


# =============================================================================
# DATA QUALITY SHEET
# =============================================================================

def create_data_quality_sheet(
    wb,
    df,
):
    """
    Create feature-level data quality analysis.
    """

    ws = wb.create_sheet(
        "Data Quality"
    )

    ws.sheet_view.showGridLines = False

    apply_title(
        ws,
        "DATA QUALITY ANALYSIS",
        (
            "Missing values, data types, uniqueness, "
            "and descriptive statistics"
        ),
        end_column=10,
    )

    rows = []

    for column in df.columns:

        series = df[column]

        missing = int(series.isna().sum())

        total = len(series)

        missing_pct = (
            missing / total * 100
            if total
            else 0
        )

        unique = int(
            series.nunique(
                dropna=True
            )
        )

        dtype = str(series.dtype)

        rows.append({
            "Feature": column,
            "Data Type": dtype,
            "Total Records": total,
            "Missing Values": missing,
            "Missing %": round(
                missing_pct,
                2,
            ),
            "Unique Values": unique,
        })

    quality_df = pd.DataFrame(rows)

    start_row = 4

    for col_index, column in enumerate(
        quality_df.columns,
        start=1,
    ):

        ws.cell(
            start_row,
            col_index,
            column,
        )

    style_header_row(
        ws,
        start_row,
    )

    for row_index, row in enumerate(
        quality_df.itertuples(index=False),
        start=start_row + 1,
    ):

        for col_index, value in enumerate(
            row,
            start=1,
        ):

            cell = ws.cell(
                row_index,
                col_index,
                value,
            )

            cell.border = border

            if col_index == 5:
                cell.number_format = "0.00"

    end_row = (
        start_row +
        len(quality_df)
    )

    add_excel_table(
        ws,
        start_row,
        end_row,
        1,
        len(quality_df.columns),
        "DataQuality",
    )

    # Conditional formatting for missing percentage
    ws.conditional_formatting.add(
        f"E{start_row + 1}:E{end_row}",
        ColorScaleRule(
            start_type="min",
            start_color=LIGHT_GREEN,
            mid_type="percentile",
            mid_value=50,
            mid_color=LIGHT_ORANGE,
            end_type="max",
            end_color=LIGHT_RED,
        ),
    )

    ws.freeze_panes = "A5"

    auto_fit_columns(ws)


# =============================================================================
# NUMERICAL STATISTICS
# =============================================================================

def create_numeric_statistics_sheet(
    wb,
    df,
):
    """
    Create descriptive statistics for numerical features.
    """

    ws = wb.create_sheet(
        "Numerical Statistics"
    )

    ws.sheet_view.showGridLines = False

    numeric_df = df.select_dtypes(
        include=np.number
    )

    if numeric_df.empty:

        apply_title(
            ws,
            "NUMERICAL STATISTICS",
            "No numerical features were detected.",
            end_column=8,
        )

        return

    stats = numeric_df.describe().T.reset_index()

    stats = stats.rename(
        columns={
            "index": "Feature"
        }
    )

    apply_title(
        ws,
        "NUMERICAL FEATURE STATISTICS",
        (
            "Descriptive statistics before model preprocessing"
        ),
        end_column=len(stats.columns),
    )

    start_row = 4

    for col_index, column in enumerate(
        stats.columns,
        start=1,
    ):

        ws.cell(
            start_row,
            col_index,
            column,
        )

    style_header_row(
        ws,
        start_row,
    )

    for row_index, row in enumerate(
        stats.itertuples(index=False),
        start=start_row + 1,
    ):

        for col_index, value in enumerate(
            row,
            start=1,
        ):

            cell = ws.cell(
                row_index,
                col_index,
                value,
            )

            cell.border = border

            if isinstance(value, (float, np.floating)):
                cell.number_format = "#,##0.00"

    end_row = (
        start_row +
        len(stats)
    )

    add_excel_table(
        ws,
        start_row,
        end_row,
        1,
        len(stats.columns),
        "NumericalStatistics",
    )

    ws.freeze_panes = "A5"

    auto_fit_columns(ws)


# =============================================================================
# TARGET ANALYSIS SHEET
# =============================================================================

def create_target_sheet(
    wb,
    df,
):
    """
    Create target distribution and class balance analysis.
    """

    ws = wb.create_sheet(
        "Target Analysis"
    )

    ws.sheet_view.showGridLines = False

    apply_title(
        ws,
        "TARGET VARIABLE ANALYSIS",
        "Income-class distribution and class balance",
        end_column=7,
    )

    if "income" not in df.columns:

        ws["A4"] = (
            "Target column 'income' was not found."
        )

        return

    target = (
        df["income"]
        .map(normalize_target)
    )

    counts = target.value_counts()

    total = counts.sum()

    rows = []

    for label, count in counts.items():

        percentage = (
            count / total * 100
            if total
            else 0
        )

        rows.append([
            label,
            int(count),
            percentage,
        ])

    headers = [
        "Income Class",
        "Records",
        "Percentage",
    ]

    start_row = 4

    for col_index, header in enumerate(
        headers,
        start=1,
    ):

        ws.cell(
            start_row,
            col_index,
            header,
        )

    style_header_row(
        ws,
        start_row,
    )

    for row_index, row in enumerate(
        rows,
        start=start_row + 1,
    ):

        for col_index, value in enumerate(
            row,
            start=1,
        ):

            cell = ws.cell(
                row_index,
                col_index,
                value,
            )

            cell.border = border

            if col_index == 3:
                cell.number_format = "0.00"

    end_row = (
        start_row +
        len(rows)
    )

    add_excel_table(
        ws,
        start_row,
        end_row,
        1,
        3,
        "TargetAnalysis",
    )

    # -------------------------------------------------------------------------
    # BAR CHART
    # -------------------------------------------------------------------------

    chart = BarChart()

    data = Reference(
        ws,
        min_col=2,
        min_row=4,
        max_row=end_row,
    )

    categories = Reference(
        ws,
        min_col=1,
        min_row=5,
        max_row=end_row,
    )

    chart.add_data(
        data,
        titles_from_data=True,
    )

    chart.set_categories(
        categories
    )

    chart.title = "Income Class Distribution"

    chart.y_axis.title = "Records"
    chart.x_axis.title = "Income Class"

    chart.height = 8
    chart.width = 13

    chart.legend = None

    ws.add_chart(
        chart,
        "E4",
    )

    ws.freeze_panes = "A5"

    auto_fit_columns(ws)


# =============================================================================
# TRAIN TEST SHEET
# =============================================================================

def create_split_sheet(
    wb,
    train_df,
    test_df,
):
    """
    Create train/test split summary.
    """

    ws = wb.create_sheet(
        "Train Test Split"
    )

    ws.sheet_view.showGridLines = False

    apply_title(
        ws,
        "TRAIN / TEST SPLIT",
        "Dataset partitioning used by Task 1",
        end_column=8,
    )

    total = len(train_df) + len(test_df)

    train_pct = (
        len(train_df) / total * 100
        if total
        else 0
    )

    test_pct = (
        len(test_df) / total * 100
        if total
        else 0
    )

    rows = [
        [
            "Training",
            len(train_df),
            train_pct,
        ],
        [
            "Testing",
            len(test_df),
            test_pct,
        ],
        [
            "Total",
            total,
            100,
        ],
    ]

    headers = [
        "Partition",
        "Records",
        "Percentage",
    ]

    start_row = 4

    for col_index, header in enumerate(
        headers,
        start=1,
    ):

        ws.cell(
            start_row,
            col_index,
            header,
        )

    style_header_row(
        ws,
        start_row,
    )

    for row_index, row in enumerate(
        rows,
        start=start_row + 1,
    ):

        for col_index, value in enumerate(
            row,
            start=1,
        ):

            cell = ws.cell(
                row_index,
                col_index,
                value,
            )

            cell.border = border

            if col_index == 3:
                cell.number_format = "0.00"

    add_excel_table(
        ws,
        start_row,
        start_row + len(rows),
        1,
        3,
        "TrainTestSplit",
    )

    # -------------------------------------------------------------------------
    # TRAINING / TESTING TARGET DISTRIBUTION
    # -------------------------------------------------------------------------

    if "income" in train_df.columns:

        ws["A10"] = "Training Target Distribution"
        ws["A10"].fill = subheader_fill
        ws["A10"].font = Font(
            color=WHITE,
            bold=True,
        )

        train_target = (
            train_df["income"]
            .map(normalize_target)
            .value_counts()
        )

        ws["A11"] = "Class"
        ws["B11"] = "Records"

        style_header_row(
            ws,
            11,
        )

        for row_number, (label, count) in enumerate(
            train_target.items(),
            start=12,
        ):

            ws.cell(
                row_number,
                1,
                label,
            )

            ws.cell(
                row_number,
                2,
                int(count),
            )

            ws.cell(
                row_number,
                1,
            ).border = border

            ws.cell(
                row_number,
                2,
            ).border = border

    if "income" in test_df.columns:

        ws["E10"] = "Testing Target Distribution"
        ws["E10"].fill = subheader_fill
        ws["E10"].font = Font(
            color=WHITE,
            bold=True,
        )

        test_target = (
            test_df["income"]
            .map(normalize_target)
            .value_counts()
        )

        ws["E11"] = "Class"
        ws["F11"] = "Records"

        for cell in ws[11][4:6]:
            cell.fill = header_fill
            cell.font = Font(
                color=WHITE,
                bold=True,
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

        for row_number, (label, count) in enumerate(
            test_target.items(),
            start=12,
        ):

            ws.cell(
                row_number,
                5,
                label,
            )

            ws.cell(
                row_number,
                6,
                int(count),
            )

            ws.cell(
                row_number,
                5,
            ).border = border

            ws.cell(
                row_number,
                6,
            ).border = border

    ws.freeze_panes = "A5"

    auto_fit_columns(ws)


# =============================================================================
# FEATURE INVENTORY
# =============================================================================

def create_feature_inventory(
    wb,
    df,
):
    """
    Create a feature inventory useful for documentation.
    """

    ws = wb.create_sheet(
        "Feature Inventory"
    )

    ws.sheet_view.showGridLines = False

    apply_title(
        ws,
        "FEATURE INVENTORY",
        "Feature names and preprocessing treatment",
        end_column=7,
    )

    numeric_features = [
        "age",
        "fnlwgt",
        "education-num",
        "capital-gain",
        "capital-loss",
        "hours-per-week",
    ]

    categorical_features = [
        "workclass",
        "education",
        "marital-status",
        "occupation",
        "relationship",
        "race",
        "sex",
        "native-country",
    ]

    rows = []

    for column in df.columns:

        if column == "income":

            rows.append([
                column,
                "Target",
                "Binary Classification",
                "Target encoding",
                "0 = <=50K; 1 = >50K",
                int(df[column].isna().sum()),
            ])

        elif column in numeric_features:

            rows.append([
                column,
                "Numerical",
                "Continuous / Integer",
                "Median → Outlier Clip → StandardScaler",
                "1st–99th percentile clipping",
                int(df[column].isna().sum()),
            ])

        elif column in categorical_features:

            rows.append([
                column,
                "Categorical",
                "Nominal",
                "Most Frequent → OneHotEncoder",
                "handle_unknown='ignore'",
                int(df[column].isna().sum()),
            ])

        else:

            rows.append([
                column,
                "Other",
                str(df[column].dtype),
                "Pipeline-dependent",
                "",
                int(df[column].isna().sum()),
            ])

    headers = [
        "Feature",
        "Type",
        "Data Role",
        "Preprocessing",
        "Details",
        "Missing Values",
    ]

    start_row = 4

    for col_index, header in enumerate(
        headers,
        start=1,
    ):

        ws.cell(
            start_row,
            col_index,
            header,
        )

    style_header_row(
        ws,
        start_row,
    )

    for row_number, row in enumerate(
        rows,
        start=start_row + 1,
    ):

        for col_number, value in enumerate(
            row,
            start=1,
        ):

            cell = ws.cell(
                row_number,
                col_number,
                value,
            )

            cell.border = border

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

    end_row = (
        start_row +
        len(rows)
    )

    add_excel_table(
        ws,
        start_row,
        end_row,
        1,
        len(headers),
        "FeatureInventory",
    )

    ws.freeze_panes = "A5"

    auto_fit_columns(ws)


# =============================================================================
# WORKBOOK METADATA
# =============================================================================

def configure_workbook(wb):
    """
    Configure workbook properties.
    """

    wb.properties.title = (
        "Adult Income ML Project - "
        "Task 1 Data Preprocessing Report"
    )

    wb.properties.subject = (
        "Professional data preprocessing report"
    )

    wb.properties.creator = (
        "Adult Income ML Project"
    )

    wb.properties.description = (
        "Automated Excel report for Task 1 "
        "data preprocessing."
    )

    wb.properties.keywords = (
        "machine learning, preprocessing, "
        "adult income, UCI, data science"
    )


# =============================================================================
# MAIN REPORT GENERATOR
# =============================================================================

def generate_report():

    print("=" * 80)
    print("CREATING PROFESSIONAL EXCEL REPORT")
    print("=" * 80)

    (
        cleaned_df,
        summary_df,
        train_df,
        test_df,
        y_train_df,
        y_test_df,
    ) = load_project_data()

    print(
        f"Cleaned dataset: {cleaned_df.shape}"
    )

    print(
        f"Training dataset: {train_df.shape}"
    )

    print(
        f"Testing dataset: {test_df.shape}"
    )

    wb = Workbook()

    # Remove default sheet
    default_sheet = wb.active

    wb.remove(default_sheet)

    configure_workbook(wb)

    # -------------------------------------------------------------------------
    # CREATE SHEETS
    # -------------------------------------------------------------------------

    print("\nCreating Dashboard...")
    create_dashboard(
        wb,
        cleaned_df,
        train_df,
        test_df,
        summary_df,
    )

    print("Creating Preprocessing Summary...")
    create_summary_sheet(
        wb,
        summary_df,
    )

    print("Creating Data Quality...")
    create_data_quality_sheet(
        wb,
        cleaned_df,
    )

    print("Creating Numerical Statistics...")
    create_numeric_statistics_sheet(
        wb,
        cleaned_df,
    )

    print("Creating Target Analysis...")
    create_target_sheet(
        wb,
        cleaned_df,
    )

    print("Creating Train/Test Split...")
    create_split_sheet(
        wb,
        train_df,
        test_df,
    )

    print("Creating Feature Inventory...")
    create_feature_inventory(
        wb,
        cleaned_df,
    )

    # -------------------------------------------------------------------------
    # TAB COLORS
    # -------------------------------------------------------------------------

    tab_colors = {
        "Dashboard": BLUE,
        "Preprocessing Summary": CYAN,
        "Data Quality": ORANGE,
        "Numerical Statistics": GREEN,
        "Target Analysis": PURPLE,
        "Train Test Split": BLUE,
        "Feature Inventory": DARK_BLUE,
    }

    for sheet_name, color in tab_colors.items():

        if sheet_name in wb.sheetnames:

            wb[sheet_name].sheet_properties.tabColor = color

    # -------------------------------------------------------------------------
    # SAVE
    # -------------------------------------------------------------------------

    wb.save(
        OUTPUT_FILE
    )

    print("\n" + "=" * 80)
    print("EXCEL REPORT CREATED SUCCESSFULLY")
    print("=" * 80)

    print(
        f"\nOutput:\n{OUTPUT_FILE}"
    )

    print(
        "\nSheets:"
    )

    for sheet_name in wb.sheetnames:

        if not sheet_name.startswith("_"):

            print(
                f"  ✓ {sheet_name}"
            )

    print("\nDone.")


# =============================================================================
# SCRIPT ENTRY POINT
# =============================================================================

if __name__ == "__main__":

    try:

        generate_report()

    except FileNotFoundError as error:

        print("\nERROR:")
        print(error)

        sys.exit(1)

    except Exception as error:

        print("\nUNEXPECTED ERROR:")
        print(
            f"{type(error).__name__}: {error}"
        )

        raise